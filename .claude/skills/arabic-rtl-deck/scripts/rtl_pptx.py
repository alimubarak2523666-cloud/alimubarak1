#!/usr/bin/env python3
"""rtl_pptx - build, repair and audit Arabic (RTL) PowerPoint decks.

PowerPoint only renders Arabic correctly when each paragraph is marked RTL
(<a:pPr rtl="1">) and each Arabic run names a complex-script font (<a:cs>).
Decks made by python-pptx, Google Slides exports or English templates usually
have neither, so Arabic bullets start on the wrong side, punctuation jumps to
the wrong end and Arabic falls back to a random font.

Usage:
    python rtl_pptx.py audit IN.pptx
    python rtl_pptx.py fix   IN.pptx OUT.pptx [--font "Cairo"]
    python rtl_pptx.py build OUTLINE.json OUT.pptx

Requires: pip install python-pptx
"""
import argparse
import json
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ARABIC = re.compile(r"[؀-ۿݐ-ݿࢠ-ࣿﭐ-﷿ﹰ-﻿]")
LATIN = re.compile(r"[A-Za-z]")
DEFAULT_AR_FONT = "Cairo"  # Ali's Arabic font
DEFAULT_LATIN_FONT = "Inter"


# --------------------------------------------------------------------------- walking

def iter_text_frames(shapes):
    """Yield every text frame in a shape tree, including groups and tables."""
    for shp in shapes:
        if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_text_frames(shp.shapes)
            continue
        if getattr(shp, "has_text_frame", False) and shp.has_text_frame:
            yield shp.text_frame
        if getattr(shp, "has_table", False) and shp.has_table:
            for row in shp.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def all_frames(prs, notes=True):
    for idx, slide in enumerate(prs.slides, 1):
        for tf in iter_text_frames(slide.shapes):
            yield idx, tf
        if notes and slide.has_notes_slide:
            yield idx, slide.notes_slide.notes_text_frame


# --------------------------------------------------------------------------- xml helpers

def _pPr(p):
    return p._p.get_or_add_pPr()


def _set_run_fonts(run_el, ar_font, latin_font=None):
    rPr = run_el.find(qn("a:rPr"))
    if rPr is None:
        rPr = run_el.makeelement(qn("a:rPr"), {})
        run_el.insert(0, rPr)
    rPr.set("lang", rPr.get("lang") or "ar-KW")
    # element order inside rPr matters: ln, fill, effect, highlight, uLnTx.., latin, ea, cs, sym
    for tag, face in (("a:latin", latin_font), ("a:cs", ar_font)):
        if not face:
            continue
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            # insert before ea/cs/sym/hlinkClick etc. to keep schema order
            anchors = [qn(t) for t in ("a:ea", "a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")]
            if tag == "a:cs":
                anchors = anchors[2:]
            pos = next((i for i, ch in enumerate(rPr) if ch.tag in anchors), len(rPr))
            rPr.insert(pos, el)
        el.set("typeface", face)
    return rPr


# --------------------------------------------------------------------------- audit

def audit(path):
    prs = Presentation(path)
    issues = []
    ar_paras = 0
    for slide_no, tf in all_frames(prs):
        for p in tf.paragraphs:
            text = "".join(r.text for r in p.runs)
            if not ARABIC.search(text):
                continue
            ar_paras += 1
            snippet = text.strip()[:40]
            pPr = p._p.pPr
            if pPr is None or pPr.get("rtl") != "1":
                issues.append((slide_no, "error", "paragraph not marked RTL", snippet))
            algn = pPr.get("algn") if pPr is not None else None
            if algn == "l":
                issues.append((slide_no, "warn", "Arabic paragraph left-aligned", snippet))
            for r in p.runs:
                if not ARABIC.search(r.text):
                    continue
                rPr = r._r.find(qn("a:rPr"))
                cs = rPr.find(qn("a:cs")) if rPr is not None else None
                if cs is None or not cs.get("typeface"):
                    issues.append((slide_no, "warn", "Arabic run has no complex-script (cs) font", r.text.strip()[:40]))
                if rPr is not None and rPr.get("spc") not in (None, "0"):
                    issues.append((slide_no, "error", "letter spacing on Arabic breaks letter joins", r.text.strip()[:40]))
            if re.search(r"[٠-٩]", text) and re.search(r"[0-9]", text):
                issues.append((slide_no, "warn", "mixes Arabic-Indic and Western digits", snippet))
    for s, lvl, msg, snip in issues:
        print(f"slide {s:>2}  {lvl:<5}  {msg}  «{snip}»")
    errs = sum(1 for i in issues if i[1] == "error")
    print(f"rtl_pptx audit: {ar_paras} Arabic paragraph(s), {errs} error(s), {len(issues) - errs} warning(s)")
    return 1 if errs else 0


# --------------------------------------------------------------------------- fix

def fix(src, dst, ar_font=DEFAULT_AR_FONT):
    prs = Presentation(src)
    changed = 0
    for _, tf in all_frames(prs):
        for p in tf.paragraphs:
            text = "".join(r.text for r in p.runs)
            if not ARABIC.search(text):
                continue
            pPr = _pPr(p)
            pPr.set("rtl", "1")
            if pPr.get("algn") in (None, "l"):
                pPr.set("algn", "r")
            for r in p.runs:
                rPr = _set_run_fonts(r._r, ar_font)
                if ARABIC.search(r.text) and rPr.get("spc"):
                    del rPr.attrib["spc"]
            changed += 1
    prs.save(dst)
    print(f"rtl_pptx fix: marked {changed} Arabic paragraph(s) RTL with cs font '{ar_font}' -> {dst}")


# --------------------------------------------------------------------------- build

W, H = Emu(12192000), Emu(6858000)  # 16:9
MARGIN = Emu(685800)


def _hex(c):
    return RGBColor.from_string(c.lstrip("#").upper())


def _para(tf, text, size, bold=False, color="1A1A1A", first=False, ar_font=DEFAULT_AR_FONT,
          latin_font=DEFAULT_LATIN_FONT, align=PP_ALIGN.RIGHT, space_after=8, bullet=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    pPr = _pPr(p)
    # Only Arabic-bearing paragraphs are RTL. A pure number/Latin line such as
    # "+38%" or "1.9x" marked RTL gets its sign moved to the wrong end.
    pPr.set("rtl", "1" if ARABIC.search(text) else "0")
    if bullet:
        pPr.set("marL", str(Emu(342900)))
        pPr.set("indent", str(-Emu(342900)))
        bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
        pPr.append(bu)
    p.space_after = Pt(space_after)
    p.line_spacing = 1.25
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = _hex(color)
    f.name = latin_font
    _set_run_fonts(r._r, ar_font, latin_font)
    return p


def _textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    bodyPr.set("rtlCol", "1")
    return tf


def _accent_bar(slide, accent):
    # accent on the RIGHT edge: the reading start in RTL
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, W - Emu(137160), 0, Emu(137160), H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = _hex(accent)
    bar.line.fill.background()


def _bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _hex(color)


def build(outline_path, dst):
    spec = json.load(open(outline_path, encoding="utf-8"))
    ar_font = spec.get("font", DEFAULT_AR_FONT)
    latin = spec.get("latin_font", DEFAULT_LATIN_FONT)
    accent = spec.get("accent", "C9A86A")
    ink = spec.get("ink", "1A1A1A")
    bg = spec.get("background", "FFFFFF")
    muted = spec.get("muted", "6B6B6B")
    kw = dict(ar_font=ar_font, latin_font=latin)

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    content_w = W - 2 * MARGIN

    # title slide
    s = prs.slides.add_slide(blank)
    _bg(s, spec.get("title_background", "0E0D11"))
    _accent_bar(s, accent)
    tf = _textbox(s, MARGIN, Emu(2286000), content_w, Emu(2286000))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _para(tf, spec["title"], 40, True, "FFFFFF", first=True, **kw)
    if spec.get("subtitle"):
        _para(tf, spec["subtitle"], 20, False, accent, **kw)
    if spec.get("presenter"):
        _para(tf, spec["presenter"], 14, False, "BFBFBF", **kw)

    for i, sl in enumerate(spec.get("slides", []), 2):
        s = prs.slides.add_slide(blank)
        kind = sl.get("type", "bullets")
        if kind == "section":
            _bg(s, spec.get("title_background", "0E0D11"))
            _accent_bar(s, accent)
            tf = _textbox(s, MARGIN, Emu(2743200), content_w, Emu(1371600))
            _para(tf, sl["title"], 36, True, "FFFFFF", first=True, **kw)
            continue
        _bg(s, bg)
        _accent_bar(s, accent)
        tf = _textbox(s, MARGIN, Emu(457200), content_w, Emu(914400))
        _para(tf, sl["title"], 30, True, ink, first=True, **kw)
        body_top = Emu(1600200)
        body_h = H - body_top - Emu(685800)
        if kind == "bullets":
            tf = _textbox(s, MARGIN, body_top, content_w, body_h)
            for j, b in enumerate(sl.get("bullets", [])):
                _para(tf, b, 20, False, ink, first=(j == 0), bullet=True, space_after=12, **kw)
        elif kind == "stats":
            stats = sl.get("stats", [])
            n = max(1, len(stats))
            gap = Emu(228600)
            col_w = int((content_w - gap * (n - 1)) / n)
            for j, st in enumerate(stats):
                # first stat on the right
                left = W - MARGIN - (j + 1) * col_w - j * gap
                tf = _textbox(s, left, body_top + Emu(457200), col_w, Emu(2286000))
                _para(tf, st["value"], 44, True, accent, first=True, align=PP_ALIGN.CENTER, **kw)
                _para(tf, st["label"], 16, False, muted, align=PP_ALIGN.CENTER, **kw)
        elif kind == "two_col":
            gap = Emu(457200)
            col_w = int((content_w - gap) / 2)
            for j, key in enumerate(("first", "second")):  # first = right column
                left = W - MARGIN - (j + 1) * col_w - j * gap
                tf = _textbox(s, left, body_top, col_w, body_h)
                col = sl.get(key, {})
                _para(tf, col.get("heading", ""), 22, True, accent, first=True, **kw)
                for b in col.get("bullets", []):
                    _para(tf, b, 18, False, ink, bullet=True, **kw)
        if sl.get("notes"):
            s.notes_slide.notes_text_frame.text = sl["notes"]
            for p in s.notes_slide.notes_text_frame.paragraphs:
                _pPr(p).set("rtl", "1")
                for r in p.runs:
                    _set_run_fonts(r._r, ar_font)
        # slide number, bottom-left (end side in RTL)
        tf = _textbox(s, MARGIN - Emu(228600), H - Emu(548640), Emu(914400), Emu(365760))
        _para(tf, str(i), 11, False, muted, first=True, align=PP_ALIGN.LEFT, **kw)

    prs.save(dst)
    print(f"rtl_pptx build: {len(prs.slides)} slide(s) -> {dst}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("audit"); a.add_argument("src")
    f = sub.add_parser("fix"); f.add_argument("src"); f.add_argument("dst"); f.add_argument("--font", default=DEFAULT_AR_FONT)
    b = sub.add_parser("build"); b.add_argument("outline"); b.add_argument("dst")
    args = ap.parse_args()
    if args.cmd == "audit":
        sys.exit(audit(args.src))
    if args.cmd == "fix":
        fix(args.src, args.dst, args.font)
    if args.cmd == "build":
        build(args.outline, args.dst)


if __name__ == "__main__":
    main()
